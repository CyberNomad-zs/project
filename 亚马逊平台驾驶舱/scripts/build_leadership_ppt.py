# -*- coding: utf-8 -*-
"""Generate leadership PPT for ASIN health project."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import nsmap
from pptx.oxml import parse_xml
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "ASIN健康度项目进展汇报.pptx"

# Brand-neutral dark blue palette
C_BG = RGBColor(0x0F, 0x17, 0x2A)
C_TITLE = RGBColor(0x0F, 0x17, 0x2A)
C_ACCENT = RGBColor(0x1D, 0x4E, 0xD8)
C_ACCENT2 = RGBColor(0x0E, 0xA5, 0xE9)
C_GREEN = RGBColor(0x05, 0x96, 0x69)
C_ORANGE = RGBColor(0xD9, 0x77, 0x06)
C_RED = RGBColor(0xDC, 0x26, 0x26)
C_MUTED = RGBColor(0x64, 0x74, 0x8B)
C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
C_CARD = RGBColor(0xF1, 0xF5, 0xF9)
C_LINE = RGBColor(0xE2, 0xE8, 0xF0)
C_BODY = RGBColor(0x1E, 0x29, 0x3B)


def set_run(run, text, size=18, bold=False, color=C_BODY, font="Microsoft YaHei"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    try:
        from pptx.oxml.ns import qn

        rPr = run._r.get_or_add_rPr()
        ea = rPr.find(qn("a:ea"))
        if ea is None:
            ea = parse_xml(
                f'<a:ea xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="{font}"/>'
            )
            rPr.append(ea)
        else:
            ea.set("typeface", font)
    except Exception:
        pass


def add_para(tf, text, size=16, bold=False, color=C_BODY, space_before=6, space_after=4, align=PP_ALIGN.LEFT):
    p0 = tf.paragraphs[0]
    if (not p0.runs) and p0.text == "" and len(tf.paragraphs) == 1:
        p = p0
    else:
        p = tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    run = p.add_run()
    set_run(run, text, size, bold, color)
    return p


def fill_shape(shape, color: RGBColor):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_rect(slide, left, top, width, height, color):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    fill_shape(sh, color)
    # softer corners
    try:
        sh.adjustments[0] = 0.08
    except Exception:
        pass
    return sh


def add_bar(slide, left, top, width, height, color):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    fill_shape(sh, color)
    return sh


def footer(slide, page, total=10):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(7.05), Inches(8.5), Inches(0.3))
    tf = box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    set_run(run, f"ASIN 健康度监控与运营闭环  |  机密·内部  |  {page}/{total}", 10, False, C_MUTED)
    p2_box = slide.shapes.add_textbox(Inches(11.2), Inches(7.05), Inches(1.5), Inches(0.3))
    p2 = p2_box.text_frame.paragraphs[0]
    p2.alignment = PP_ALIGN.RIGHT
    r2 = p2.add_run()
    set_run(r2, "2026-07-22", 10, False, C_MUTED)


def title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # left accent panel
    add_bar(slide, Inches(0), Inches(0), Inches(0.18), Inches(7.5), C_ACCENT)
    add_bar(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.08), C_ACCENT)

    box = slide.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.5), Inches(1))
    tf = box.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    set_run(r, "亚马逊 ASIN 健康度监控与运营闭环", 32, True, C_TITLE)

    box2 = slide.shapes.add_textbox(Inches(0.8), Inches(3.0), Inches(11.5), Inches(0.6))
    tf2 = box2.text_frame
    p2 = tf2.paragraphs[0]
    r2 = p2.add_run()
    set_run(r2, "项目进展汇报  ·  领导汇报版", 20, False, C_ACCENT)

    box3 = slide.shapes.add_textbox(Inches(0.8), Inches(4.0), Inches(11), Inches(1.2))
    tf3 = box3.text_frame
    add_para(tf3, "赛狐数据自动同步 → 规则识别异常 → 飞书任务闭环 → 群机器人提醒", 16, False, C_BODY)
    add_para(tf3, "一期主链路已打通，进入业务试运行阶段", 16, True, C_GREEN)
    add_para(tf3, "汇报日期：2026-07-22    样本数据日：2026-07-21", 14, False, C_MUTED)

    # bottom status chips
    chips = [
        (Inches(0.8), "技术主路径 ~80%", C_ACCENT),
        (Inches(3.6), "日同步 110 ASIN", C_GREEN),
        (Inches(6.4), "待办任务 28 条", C_ORANGE),
        (Inches(9.2), "日批 10:00 已通", C_ACCENT2),
    ]
    for left, text, color in chips:
        card = add_rect(slide, left, Inches(5.8), Inches(2.5), Inches(0.55), color)
        tb = slide.shapes.add_textbox(left, Inches(5.9), Inches(2.5), Inches(0.4))
        p = tb.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        set_run(r, text, 12, True, C_WHITE)


def section_header(slide, title, subtitle=None):
    add_bar(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.08), C_ACCENT)
    add_bar(slide, Inches(0.5), Inches(0.35), Inches(0.12), Inches(0.45), C_ACCENT)
    box = slide.shapes.add_textbox(Inches(0.8), Inches(0.3), Inches(11), Inches(0.55))
    p = box.text_frame.paragraphs[0]
    r = p.add_run()
    set_run(r, title, 24, True, C_TITLE)
    if subtitle:
        box2 = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.5), Inches(0.35))
        p2 = box2.text_frame.paragraphs[0]
        r2 = p2.add_run()
        set_run(r2, subtitle, 12, False, C_MUTED)


def slide_conclusion(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    section_header(slide, "01  一句话结论", "给领导的核心判断")
    footer(slide, 2)

    # big quote card
    add_rect(slide, Inches(0.6), Inches(1.5), Inches(12.1), Inches(2.0), C_CARD)
    box = slide.shapes.add_textbox(Inches(0.9), Inches(1.7), Inches(11.5), Inches(1.6))
    tf = box.text_frame
    tf.word_wrap = True
    add_para(
        tf,
        "已实现：赛狐经营数据每日自动同步 → 规则识别黄/橙/红异常 → 任务落到飞书 → 群机器人推送摘要。",
        18,
        True,
        C_TITLE,
        4,
        8,
    )
    add_para(
        tf,
        "老板与负责人只需看重点异常和待决策事项，无需再人工从赛狐抄数、拼表。",
        16,
        False,
        C_BODY,
        4,
        4,
    )

    # three points
    items = [
        ("技术", "一期主路径约 80%\n日同步 / 规则 / 飞书 / 机器人 / 定时", C_ACCENT),
        ("业务", "样本日 6 店 · 110 ASIN\n待办 28（红15 / 橙13）", C_ORANGE),
        ("阶段", "进入 2～4 周试运行\n校准阈值 · 养成闭环习惯", C_GREEN),
    ]
    for i, (h, body, color) in enumerate(items):
        left = Inches(0.6 + i * 4.15)
        add_rect(slide, left, Inches(3.9), Inches(3.95), Inches(2.5), C_CARD)
        add_bar(slide, left, Inches(3.9), Inches(3.95), Inches(0.12), color)
        tb = slide.shapes.add_textbox(left + Inches(0.2), Inches(4.2), Inches(3.5), Inches(2.0))
        tf = tb.text_frame
        tf.word_wrap = True
        add_para(tf, h, 16, True, color, 0, 8)
        for line in body.split("\n"):
            add_para(tf, line, 14, False, C_BODY, 2, 2)


def slide_background(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    section_header(slide, "02  背景与目标", "不重做赛狐分析 · 不做利润 · 不做自动操作后台")
    footer(slide, 3)

    roles = [
        ("赛狐", "数据发生了什么变化", "销售 / 流量 / 广告 / 库存 / 评价"),
        ("飞书", "谁负责、做了什么、何时复查", "任务 · 动作 · 闭环"),
        ("老板", "只看重点与待决策", "红任务 · 逾期 · 需拍板事项"),
    ]
    for i, (t, s, d) in enumerate(roles):
        left = Inches(0.6 + i * 4.15)
        add_rect(slide, left, Inches(1.5), Inches(3.95), Inches(2.4), C_CARD)
        add_bar(slide, left, Inches(1.5), Inches(0.12), Inches(2.4), C_ACCENT)
        tb = slide.shapes.add_textbox(left + Inches(0.3), Inches(1.7), Inches(3.4), Inches(2.0))
        tf = tb.text_frame
        tf.word_wrap = True
        add_para(tf, t, 20, True, C_ACCENT, 0, 8)
        add_para(tf, s, 15, True, C_TITLE, 4, 6)
        add_para(tf, d, 13, False, C_MUTED, 2, 2)

    # goals
    add_rect(slide, Inches(0.6), Inches(4.2), Inches(12.1), Inches(2.4), C_CARD)
    tb = slide.shapes.add_textbox(Inches(0.9), Inches(4.35), Inches(11.5), Inches(2.1))
    tf = tb.text_frame
    tf.word_wrap = True
    add_para(tf, "一期五件事", 16, True, C_TITLE, 0, 6)
    goals = [
        "1）每日自动同步赛狐 ASIN 数据至飞书",
        "2）按可配置规则生成橙/红异常任务并匹配负责人",
        "3）防重复建单 + 群机器人摘要提醒",
        "4）运营可记录动作与复查（表结构已就绪）",
        "5）老板可筛查异常视图（红 / 负责人 / 待决策）",
    ]
    for g in goals:
        add_para(tf, g, 13, False, C_BODY, 2, 1)


def slide_architecture(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    section_header(slide, "03  系统架构与日批节奏", "每天 10:00 自动跑通全流程")
    footer(slide, 4)

    steps = [
        ("① 拉数", "赛狐 OpenAPI\n产品分析 + FBA库存"),
        ("② 落库", "本地 SQLite\n同步日志 / 幂等"),
        ("③ 规则", "黄 / 橙 / 红\n防重复建任务"),
        ("④ 飞书", "每日数据表\n异常任务表"),
        ("⑤ 提醒", "群机器人摘要\n按负责人汇总"),
    ]
    for i, (h, b) in enumerate(steps):
        left = Inches(0.4 + i * 2.55)
        add_rect(slide, left, Inches(1.6), Inches(2.35), Inches(2.3), C_CARD)
        add_bar(slide, left, Inches(1.6), Inches(2.35), Inches(0.1), C_ACCENT if i < 4 else C_GREEN)
        tb = slide.shapes.add_textbox(left + Inches(0.12), Inches(1.85), Inches(2.1), Inches(1.9))
        tf = tb.text_frame
        tf.word_wrap = True
        add_para(tf, h, 16, True, C_ACCENT, 0, 8)
        for line in b.split("\n"):
            add_para(tf, line, 12, False, C_BODY, 2, 1)
        if i < 4:
            # arrow-ish bar
            ar = slide.shapes.add_textbox(left + Inches(2.25), Inches(2.4), Inches(0.35), Inches(0.4))
            p = ar.text_frame.paragraphs[0]
            r = p.add_run()
            set_run(r, "→", 18, True, C_MUTED)

    # timeline
    add_rect(slide, Inches(0.6), Inches(4.3), Inches(12.1), Inches(2.3), C_CARD)
    tb = slide.shapes.add_textbox(Inches(0.9), Inches(4.45), Inches(11.5), Inches(2.0))
    tf = tb.text_frame
    tf.word_wrap = True
    add_para(tf, "建议日节奏", 15, True, C_TITLE, 0, 6)
    add_para(tf, "10:00  系统自动同步 + 生成任务 + 推送群摘要", 13, False, C_BODY, 2, 2)
    add_para(tf, "上午    负责人/助理优先处理红色，再过橙色", 13, False, C_BODY, 2, 2)
    add_para(tf, "下班前  助理在「运营动作记录」补齐原因、动作、复查日", 13, False, C_BODY, 2, 2)
    add_para(tf, "原则    系统只监控提醒，不自动改价/关广告/操作亚马逊后台", 13, True, C_ORANGE, 4, 2)


def slide_done(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    section_header(slide, "04  已交付成果", "可演示 · 可试用")
    footer(slide, 5)

    left_items = [
        ("✅", "赛狐 OpenAPI", "鉴权、签名、限流、产品分析 V2、FBA 库存"),
        ("✅", "日数据同步", "真网验证，样本日 110 个 ASIN"),
        ("✅", "规则引擎", "多指标黄/橙/红 + 防重复"),
        ("✅", "飞书四表", "每日数据 / 任务 / 动作 / 配置"),
        ("✅", "机器人提醒", "群摘要推送已实测成功"),
        ("✅", "Windows 日批", "每天 10:00，已修复并验证成功"),
    ]
    add_rect(slide, Inches(0.5), Inches(1.4), Inches(7.4), Inches(5.3), C_CARD)
    tb = slide.shapes.add_textbox(Inches(0.75), Inches(1.55), Inches(7.0), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True
    for mark, title, desc in left_items:
        add_para(tf, f"{mark}  {title}", 15, True, C_TITLE, 6, 2)
        add_para(tf, f"     {desc}", 12, False, C_MUTED, 0, 4)

    # right status
    add_rect(slide, Inches(8.15), Inches(1.4), Inches(4.6), Inches(2.4), C_ACCENT)
    tb = slide.shapes.add_textbox(Inches(8.4), Inches(1.6), Inches(4.2), Inches(2.0))
    tf = tb.text_frame
    tf.word_wrap = True
    add_para(tf, "当前阶段", 14, False, C_WHITE, 0, 6)
    add_para(tf, "业务试运行", 26, True, C_WHITE, 2, 8)
    add_para(tf, "技术主路径已通\n阈值与动作习惯校准中", 13, False, C_WHITE, 2, 2)

    add_rect(slide, Inches(8.15), Inches(4.05), Inches(4.6), Inches(2.65), C_CARD)
    tb = slide.shapes.add_textbox(Inches(8.4), Inches(4.2), Inches(4.2), Inches(2.4))
    tf = tb.text_frame
    tf.word_wrap = True
    add_para(tf, "覆盖范围（样本）", 14, True, C_TITLE, 0, 8)
    add_para(tf, "· 特米 / Ledvia → 潘启豪", 12, False, C_BODY, 2, 2)
    add_para(tf, "· 昇新 / 起的快 → 马文静", 12, False, C_BODY, 2, 2)
    add_para(tf, "· Nottac / 字同道合 → 梦晴", 12, False, C_BODY, 2, 2)
    add_para(tf, "· 共 6 店 · 110 ASIN", 12, True, C_ACCENT, 6, 2)


def slide_metrics(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    section_header(slide, "05  试运行数据快照", "业务日 2026-07-21 · 规则已多轮收紧")
    footer(slide, 6)

    kpis = [
        ("110", "同步 ASIN", C_ACCENT),
        ("28", "未关闭任务", C_ORANGE),
        ("15", "红色 · 优先", C_RED),
        ("13", "橙色 · 1 工作日", C_ORANGE),
    ]
    for i, (num, label, color) in enumerate(kpis):
        left = Inches(0.55 + i * 3.15)
        add_rect(slide, left, Inches(1.45), Inches(3.0), Inches(1.7), C_CARD)
        add_bar(slide, left, Inches(1.45), Inches(3.0), Inches(0.1), color)
        tb = slide.shapes.add_textbox(left, Inches(1.7), Inches(3.0), Inches(1.3))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        set_run(r, num, 36, True, color)
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        set_run(r2, label, 13, False, C_MUTED)

    # funnel
    add_rect(slide, Inches(0.55), Inches(3.45), Inches(6.2), Inches(3.2), C_CARD)
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(3.6), Inches(5.7), Inches(2.9))
    tf = tb.text_frame
    tf.word_wrap = True
    add_para(tf, "任务收敛过程（说明“不是一上来就刷屏”）", 14, True, C_TITLE, 0, 10)
    add_para(tf, "首轮全开规则          ≈ 110 条", 14, False, C_BODY, 4, 4)
    add_para(tf, "一轮降噪后            ≈  38 条", 14, False, C_BODY, 4, 4)
    add_para(tf, "当前试运行口径        =  28 条", 16, True, C_GREEN, 4, 8)
    add_para(tf, "已提高：库存过高门槛、ACOS 目标、无动销过滤", 12, False, C_MUTED, 4, 2)

    # structure
    add_rect(slide, Inches(7.0), Inches(3.45), Inches(5.7), Inches(3.2), C_CARD)
    tb = slide.shapes.add_textbox(Inches(7.25), Inches(3.6), Inches(5.3), Inches(2.9))
    tf = tb.text_frame
    tf.word_wrap = True
    add_para(tf, "当前任务结构", 14, True, C_TITLE, 0, 8)
    rows = [
        ("库存不足", "红 13 / 橙 3", "需核实是否真断货"),
        ("库存过高", "橙 7", "已收紧至 >240 天"),
        ("ACOS 异常", "红 1 / 橙 3", "目标暂 35%"),
        ("评价异常", "红 1", "少量关注"),
    ]
    for a, b, c in rows:
        add_para(tf, f"· {a}  {b}  — {c}", 12, False, C_BODY, 3, 2)
    add_para(tf, "负责人：梦晴 12 / 潘启豪 10 / 马文静 6", 12, True, C_ACCENT, 8, 2)


def slide_value(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    section_header(slide, "06  业务价值", "以前 vs 现在")
    footer(slide, 7)

    # table-like two columns
    add_rect(slide, Inches(0.5), Inches(1.4), Inches(6.0), Inches(5.3), C_CARD)
    add_bar(slide, Inches(0.5), Inches(1.4), Inches(6.0), Inches(0.55), RGBColor(0x94, 0xA3, 0xB8))
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(6.0), Inches(0.4))
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    set_run(r, "以前", 16, True, C_WHITE)

    tb = slide.shapes.add_textbox(Inches(0.75), Inches(2.2), Inches(5.5), Inches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True
    for t in [
        "助理从赛狐/后台抄数填 Excel",
        "异常靠人眼扫，容易漏",
        "责任边界不清",
        "老板看明细负担重",
        "处理过程难追溯",
    ]:
        add_para(tf, f"·  {t}", 14, False, C_BODY, 8, 4)

    add_rect(slide, Inches(6.8), Inches(1.4), Inches(6.0), Inches(5.3), C_CARD)
    add_bar(slide, Inches(6.8), Inches(1.4), Inches(6.0), Inches(0.55), C_GREEN)
    tb = slide.shapes.add_textbox(Inches(6.8), Inches(1.5), Inches(6.0), Inches(0.4))
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    set_run(r, "现在", 16, True, C_WHITE)

    tb = slide.shapes.add_textbox(Inches(7.05), Inches(2.2), Inches(5.5), Inches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True
    for t in [
        "系统每日自动同步，减少抄数差错",
        "规则自动标黄/橙/红，红优先",
        "任务绑定店铺负责人",
        "老板看红任务 / 待决策 / 群摘要",
        "飞书动作表可追溯闭环",
    ]:
        add_para(tf, f"·  {t}", 14, False, C_BODY, 8, 4)


def slide_risk_next(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    section_header(slide, "07  风险与下一步", "需领导知情与支持")
    footer(slide, 8)

    add_rect(slide, Inches(0.5), Inches(1.4), Inches(6.2), Inches(5.3), C_CARD)
    tb = slide.shapes.add_textbox(Inches(0.75), Inches(1.55), Inches(5.8), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True
    add_para(tf, "主要风险", 16, True, C_RED, 0, 10)
    risks = [
        ("库存红仍偏多", "业务过会核实 13 条；可再收阈值"),
        ("S/A 等级未齐", "从产品档案导入后核心款规则才准"),
        ("本机日批依赖开机登录", "长期建议固定 IP 服务器"),
        ("飞书协作者未加全", "管理员完成分享即可"),
    ]
    for a, b in risks:
        add_para(tf, f"· {a}", 13, True, C_TITLE, 6, 2)
        add_para(tf, f"  {b}", 12, False, C_MUTED, 0, 4)

    add_rect(slide, Inches(6.95), Inches(1.4), Inches(5.8), Inches(5.3), C_CARD)
    tb = slide.shapes.add_textbox(Inches(7.2), Inches(1.55), Inches(5.4), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True
    add_para(tf, "下一步 2～4 周", 16, True, C_ACCENT, 0, 10)
    steps = [
        "第 1 周：加协作者 + 红任务过会 + 动作表开填",
        "观察每日 10:00 自动任务是否稳定",
        "按误报率微调库存 / ACOS 阈值",
        "导入产品等级 S/A/B",
        "评估是否迁固定服务器",
        "一期稳定后排期二期：AI 日报等",
    ]
    for i, s in enumerate(steps, 1):
        add_para(tf, f"{i}. {s}", 13, False, C_BODY, 6, 3)


def slide_ask(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    section_header(slide, "08  需领导支持", "拍板与资源")
    footer(slide, 9)

    asks = [
        ("01", "确认试运行启动", "本周起以飞书异常任务 + 群摘要作为日常异常入口"),
        ("02", "明确过会节奏", "建议红任务日清/隔日清；周会看闭环与有效率"),
        ("03", "指定业务 Owner", "阈值争议、产品等级、扩店范围的拍板人"),
        ("04", "服务器与白名单", "若需 7×24 稳定运行，支持迁固定 IP 主机"),
        ("05", "飞书权限", "安排管理员完成表格分享与协作权限"),
    ]
    for i, (n, t, d) in enumerate(asks):
        top = Inches(1.35 + i * 1.0)
        add_rect(slide, Inches(0.6), top, Inches(12.1), Inches(0.9), C_CARD)
        # number circle-ish
        num = add_rect(slide, Inches(0.8), top + Inches(0.18), Inches(0.55), Inches(0.55), C_ACCENT)
        tb = slide.shapes.add_textbox(Inches(0.8), top + Inches(0.28), Inches(0.55), Inches(0.4))
        p = tb.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        set_run(r, n, 11, True, C_WHITE)
        tb2 = slide.shapes.add_textbox(Inches(1.55), top + Inches(0.15), Inches(10.8), Inches(0.65))
        tf = tb2.text_frame
        tf.word_wrap = True
        add_para(tf, t, 15, True, C_TITLE, 0, 2)
        add_para(tf, d, 12, False, C_MUTED, 0, 0)


def slide_end(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bar(slide, Inches(0), Inches(0), Inches(0.18), Inches(7.5), C_ACCENT)
    add_bar(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.08), C_ACCENT)

    box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.5), Inches(0.7))
    p = box.text_frame.paragraphs[0]
    r = p.add_run()
    set_run(r, "汇报三句话", 28, True, C_TITLE)

    cards = [
        "已经能自动从赛狐拉数、识别异常、写到飞书并推群消息，抄数式日报可以逐步退出。",
        "样本日 6 店 110 个 ASIN，收敛后 28 条待办（红 15 / 橙 13），可以开始业务试运行。",
        "接下来 2～4 周重点是「人用起来 + 阈值定稿」；技术侧保障每天 10:00 稳定跑通。",
    ]
    for i, text in enumerate(cards):
        top = Inches(2.4 + i * 1.15)
        add_rect(slide, Inches(0.8), top, Inches(11.7), Inches(1.0), C_CARD)
        add_bar(slide, Inches(0.8), top, Inches(0.12), Inches(1.0), C_ACCENT)
        tb = slide.shapes.add_textbox(Inches(1.2), top + Inches(0.2), Inches(11.0), Inches(0.7))
        tf = tb.text_frame
        tf.word_wrap = True
        add_para(tf, f"{i+1}.  {text}", 14, False, C_BODY, 0, 0)

    box = slide.shapes.add_textbox(Inches(0.8), Inches(6.2), Inches(11.5), Inches(0.6))
    tf = box.text_frame
    add_para(tf, "飞书工作台：https://dcnkj1up67d5.feishu.cn/base/FrXvbTodVad1Kqs0hCkckrI8noc", 11, False, C_MUTED, 0, 2)
    add_para(tf, "详细文字版：docs/项目进展汇报-ASIN健康度监控与运营闭环.md", 11, False, C_MUTED, 0, 0)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    title_slide(prs)
    slide_conclusion(prs)
    slide_background(prs)
    slide_architecture(prs)
    slide_done(prs)
    slide_metrics(prs)
    slide_value(prs)
    slide_risk_next(prs)
    slide_ask(prs)
    slide_end(prs)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"saved: {OUT}")
    print(f"slides: {len(prs.slides)}")


if __name__ == "__main__":
    main()
